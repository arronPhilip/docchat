"""DocChat - upload documents and chat with them.

Pick an AI provider with the PROVIDER setting (in .env or the environment):
  gemini  - Google Gemini, free tier (needs GEMINI_API_KEY from https://aistudio.google.com/apikey)
  ollama  - a local model via Ollama, free and offline (needs Ollama installed + a pulled model)
  claude  - Anthropic Claude, paid (needs ANTHROPIC_API_KEY); best answers, with exact citations
If PROVIDER is unset, it uses whichever key it finds (Gemini, then Claude), otherwise Ollama.

Chats and documents are saved in docchat.db next to this file.
Run:  python app.py   then open http://127.0.0.1:5000
"""

import base64
import io
import json
import os
import re
import sqlite3
import time
import uuid

from flask import Flask, Response, g, jsonify, redirect, request, send_from_directory, session, stream_with_context
from werkzeug.middleware.proxy_fix import ProxyFix

BASE_DIR = os.path.dirname(os.path.abspath(__file__))


def load_dotenv(path=os.path.join(BASE_DIR, ".env")):
    """Read KEY=value lines from .env into the environment (real environment variables win)."""
    if not os.path.exists(path):
        return
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            value = value.strip().strip('"').strip("'")
            if value:
                os.environ.setdefault(key.strip(), value)


load_dotenv()

GEMINI_MODEL = os.environ.get("GEMINI_MODEL", "gemini-3.5-flash-lite")
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "llama3.2:3b")
OLLAMA_CONTEXT = int(os.environ.get("OLLAMA_CONTEXT", "16384"))  # tokens; Ollama's default is too small for documents
CLAUDE_MODEL = "claude-opus-5-5"
DB_PATH = os.environ.get("DOCCHAT_DB", os.path.join(BASE_DIR, "docchat.db"))

# Public mode (set DOCCHAT_PUBLIC=1 on the web host): smaller uploads and daily question limits,
# so strangers can't use up the owner's free AI quota. Limits count questions + overviews; 0 = unlimited.
PUBLIC = os.environ.get("DOCCHAT_PUBLIC", "").strip() == "1"
DAILY_LIMIT = int(os.environ.get("DAILY_LIMIT", "0"))                # per visitor (IP address); 0 = unlimited
GLOBAL_DAILY_LIMIT = int(os.environ.get("GLOBAL_DAILY_LIMIT", "0"))  # whole site; 0 = unlimited

MAX_UPLOAD_BYTES = (10 if PUBLIC else 30) * 1024 * 1024  # per chat

# Sign in with Google: set GOOGLE_CLIENT_ID (from Google Cloud > APIs & Services > Credentials) to require login.
# Without it, DocChat runs without accounts (handy on your own computer).
GOOGLE_CLIENT_ID = os.environ.get("GOOGLE_CLIENT_ID", "").strip()
AUTH_ENABLED = bool(GOOGLE_CLIENT_ID)


def load_secret_key():
    """Signs the login cookie. On the host it comes from SECRET_KEY; locally it's kept in .secret_key."""
    key = os.environ.get("SECRET_KEY", "").strip()
    if key:
        return key
    path = os.path.join(BASE_DIR, ".secret_key")
    if not os.path.exists(path):
        with open(path, "w") as f:
            f.write(uuid.uuid4().hex + uuid.uuid4().hex)
    return open(path).read().strip()
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".py", ".html", ".log"}
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}

FOLLOWUP_MARK = "[[FOLLOWUPS]]"
SYSTEM_PROMPT = f"""You are DocChat, an assistant that answers questions about the documents the user uploaded.

- Base your answers on the documents. Quote the relevant parts.
- If the documents don't contain the answer, say so plainly, then offer what general knowledge you have, clearly labelled as not coming from the documents.
- Be concise unless asked otherwise. Use short paragraphs or bullet lists when that makes the answer easier to read.
- Use **bold** for key terms. Don't use Markdown headings bigger than ###.
- After your answer, write the line {FOLLOWUP_MARK} and then exactly three short follow-up questions the user might want to ask next about the documents, one per line, without numbering."""

# Gemini and local models have no built-in citations, so ask for them in the prompt
CITE_HINT = "\n- When you use a fact from a document, name the document (and page if known) in brackets, like [report.pdf, p. 3]."

STYLE_HINTS = {
    "concise": "Keep the answer short: 2-4 sentences or a few bullets.",
    "detailed": "Give a thorough, well-structured answer that covers all relevant details.",
    "simple": "Explain in simple, plain language, as if to a beginner. Avoid jargon.",
}

INSIGHTS_PROMPT = """Analyze the document(s) above. Reply with only a JSON object, no other text, in this shape:
{"title": "a short title for these documents, at most 6 words",
 "type": "what kind of document this is, 1-3 words, e.g. Contract, Research paper, Invoice, Notes",
 "summary": "a 2-3 sentence plain-language summary",
 "topics": ["4 to 6 key topics, 1-3 words each"],
 "questions": ["4 specific, useful questions a reader might ask about these documents"]}"""


def pick_provider():
    p = os.environ.get("PROVIDER", "").strip().lower()
    if p:
        return p
    if os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY"):
        return "gemini"
    if os.environ.get("ANTHROPIC_API_KEY"):
        return "claude"
    return "ollama"


PROVIDER = pick_provider()
PROVIDER_LABEL = {
    "gemini": f"Gemini · {GEMINI_MODEL}",
    "ollama": f"Ollama · {OLLAMA_MODEL} (local)",
    "claude": "Claude Opus 5.5",
}.get(PROVIDER, PROVIDER)

app = Flask(__name__, static_folder="static")
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES + 1024 * 1024
app.secret_key = load_secret_key()
app.config.update(SESSION_COOKIE_SAMESITE="Lax", SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SECURE=PUBLIC,
                  PERMANENT_SESSION_LIFETIME=60 * 60 * 24 * 30)
if PUBLIC:
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1)  # behind the host's proxy: trust its client IP

OWNER_COOKIE = "docchat_uid"


@app.before_request
def identify_visitor():
    """Signed-in users own chats by Google account; guests get a random private ID per browser."""
    g.user = session.get("user")
    uid = request.cookies.get(OWNER_COOKIE, "")
    g.new_owner = not re.fullmatch(r"[0-9a-f]{32}", uid)
    g.guest = uuid.uuid4().hex if g.new_owner else uid
    g.owner = g.user["id"] if g.user else g.guest
    if AUTH_ENABLED and not g.user and request.path.startswith("/api/chats"):
        return jsonify({"error": "Please sign in with Google first.", "auth": True}), 401


@app.after_request
def finish_response(resp):
    if getattr(g, "new_owner", False):
        resp.set_cookie(OWNER_COOKIE, g.guest, max_age=60 * 60 * 24 * 365, httponly=True,
                        samesite="Lax", secure=request.is_secure)
    # Make the browser re-check the page and its scripts every time, so an update is never hidden by an old cached copy
    if request.path in ("/", "/login", "/privacy") or request.path.startswith("/static/"):
        resp.headers["Cache-Control"] = "no-cache"
    return resp


class UserFacingError(Exception):
    """An error whose message is safe and useful to show in the chat."""


class QuotaError(UserFacingError):
    """A daily free-question limit was reached."""


# ================================================================ database

def db():
    con = sqlite3.connect(DB_PATH, timeout=10)
    con.row_factory = sqlite3.Row
    return con


def init_db():
    with db() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS chats (
            id TEXT PRIMARY KEY, title TEXT NOT NULL DEFAULT '', insights TEXT,
            created REAL NOT NULL, updated REAL NOT NULL);
        CREATE TABLE IF NOT EXISTS docs (
            id TEXT PRIMARY KEY, chat_id TEXT NOT NULL REFERENCES chats(id) ON DELETE CASCADE,
            name TEXT NOT NULL, kind TEXT NOT NULL, mime TEXT NOT NULL, size INTEGER NOT NULL,
            raw BLOB, text TEXT, meta TEXT NOT NULL DEFAULT '{}',
            turn INTEGER NOT NULL,      -- index of the user message this doc was sent with
            created REAL NOT NULL);
        CREATE TABLE IF NOT EXISTS messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            chat_id TEXT NOT NULL REFERENCES chats(id) ON DELETE CASCADE,
            role TEXT NOT NULL,         -- 'user' or 'assistant'
            text TEXT NOT NULL,         -- exactly what the model saw / said (used to rebuild history)
            display TEXT NOT NULL,      -- what the page shows
            raw TEXT,                   -- Claude only: full assistant content blocks, replayed unchanged
            meta TEXT NOT NULL DEFAULT '{}',
            created REAL NOT NULL);
        CREATE INDEX IF NOT EXISTS idx_docs_chat ON docs(chat_id);
        CREATE INDEX IF NOT EXISTS idx_messages_chat ON messages(chat_id, id);
        CREATE TABLE IF NOT EXISTS usage (key TEXT NOT NULL, day TEXT NOT NULL, count INTEGER NOT NULL,
            PRIMARY KEY (key, day));
        CREATE TABLE IF NOT EXISTS users (
            id TEXT PRIMARY KEY,        -- 'g:' + Google account ID
            email TEXT, name TEXT, picture TEXT, created REAL NOT NULL, last_login REAL NOT NULL);
        """)
        if "owner" not in [r["name"] for r in con.execute("PRAGMA table_info(chats)")]:
            con.execute("ALTER TABLE chats ADD COLUMN owner TEXT")  # databases from before per-visitor privacy
        con.execute("CREATE INDEX IF NOT EXISTS idx_chats_owner ON chats(owner, updated)")


def get_chat(con, cid):
    """The chat, only if it belongs to the current visitor."""
    return con.execute("SELECT * FROM chats WHERE id = ? AND owner = ?", (cid, g.owner)).fetchone()


def client_ip():
    return request.remote_addr or "unknown"


def today(offset_days=0):
    return time.strftime("%Y-%m-%d", time.gmtime(time.time() + offset_days * 86400))


def usage_today(con, key):
    row = con.execute("SELECT count FROM usage WHERE key = ? AND day = ?", (key, today())).fetchone()
    return row[0] if row else 0


def take_quota():
    """Count one AI request against today's limits, or raise UserFacingError if a limit is reached."""
    if not (DAILY_LIMIT or GLOBAL_DAILY_LIMIT):
        return
    ip_key = "ip:" + client_ip()
    with db() as con:
        if GLOBAL_DAILY_LIMIT and usage_today(con, "site") >= GLOBAL_DAILY_LIMIT:
            raise QuotaError("DocChat has reached its free daily limit for everyone. Please come back tomorrow.")
        if DAILY_LIMIT and usage_today(con, ip_key) >= DAILY_LIMIT:
            raise QuotaError(f"You've used all {DAILY_LIMIT} free questions for today. Please come back tomorrow.")
        for key in ("site", ip_key):
            con.execute("INSERT INTO usage (key, day, count) VALUES (?, ?, 1) "
                        "ON CONFLICT(key, day) DO UPDATE SET count = count + 1", (key, today()))
        con.execute("DELETE FROM usage WHERE day < ?", (today(-3),))


def load_docs(con, cid, with_data=False):
    cols = "*" if with_data else "id, name, kind, mime, size, meta, turn, created"
    return [dict(r) for r in con.execute(f"SELECT {cols} FROM docs WHERE chat_id = ? ORDER BY created", (cid,))]


def load_messages(con, cid):
    return [dict(r) for r in con.execute("SELECT * FROM messages WHERE chat_id = ? ORDER BY id", (cid,))]


def doc_public(d):
    meta = json.loads(d.get("meta") or "{}")
    return {"id": d["id"], "name": d["name"], "kind": d["kind"], "size": d["size"],
            "pages": meta.get("pages"), "words": meta.get("words")}


def chat_summary(con, row):
    counts = con.execute(
        "SELECT (SELECT COUNT(*) FROM docs WHERE chat_id = ?), "
        "(SELECT COUNT(*) FROM messages WHERE chat_id = ? AND role = 'user')", (row["id"], row["id"])).fetchone()
    return {"id": row["id"], "title": row["title"] or "New chat", "updated": row["updated"],
            "docs": counts[0], "questions": counts[1]}


def touch(con, cid):
    con.execute("UPDATE chats SET updated = ? WHERE id = ?", (time.time(), cid))


# ================================================================ documents

def sse(event: dict) -> str:
    return f"data: {json.dumps(event)}\n\n"


def pdf_to_text(raw: bytes):
    """Extract text page by page, with page markers so the model can cite pages."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(raw))
    pages = [f"--- Page {i} ---\n{(p.extract_text() or '').strip()}" for i, p in enumerate(reader.pages, 1)]
    return "\n\n".join(pages), len(reader.pages)


def docx_to_text(raw: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def word_count(text: str) -> int:
    return len(re.findall(r"\w+", text or ""))


def ingest(name: str, raw: bytes) -> dict:
    """Turn an uploaded file into a doc record, or raise UserFacingError."""
    ext = os.path.splitext(name)[1].lower()
    meta = {}
    if ext == ".pdf":
        try:
            text, meta["pages"] = pdf_to_text(raw)
        except Exception:
            raise UserFacingError(f"{name}: couldn't read this PDF.")
        plain = re.sub(r"--- Page \d+ ---", "", text)
        meta["words"] = word_count(plain)
        if PROVIDER == "ollama" and not plain.strip():
            raise UserFacingError(f"{name}: no text found (scanned PDF?). The local model needs text-based PDFs.")
        return {"kind": "pdf", "mime": "application/pdf", "raw": raw, "text": text, "meta": meta}
    if ext == ".docx":
        try:
            text = docx_to_text(raw)
        except Exception:
            raise UserFacingError(f"{name}: couldn't read this Word file.")
        meta["words"] = word_count(text)
        return {"kind": "text", "mime": "text/plain", "raw": None, "text": text, "meta": meta}
    if ext in IMAGE_TYPES:
        if PROVIDER == "ollama":
            raise UserFacingError(f"{name}: images need Gemini or Claude; the local model reads text only.")
        return {"kind": "image", "mime": IMAGE_TYPES[ext], "raw": raw, "text": None, "meta": meta}
    if ext in TEXT_EXTENSIONS:
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            raise UserFacingError(f"{name}: not valid UTF-8 text.")
        meta["words"] = word_count(text)
        return {"kind": "text", "mime": "text/plain", "raw": None, "text": text, "meta": meta}
    raise UserFacingError(f"{name}: unsupported type. Use PDF, Word, text or an image.")


def docs_as_text(docs) -> str:
    return "\n\n".join(f'<document name="{d["name"]}">\n{d["text"] or ""}\n</document>'
                       for d in docs if d["kind"] != "image")


def build_turns(con, cid, docs, new_user_text):
    """History as a list of turns: each user message with the docs sent alongside it, plus the reply."""
    msgs = load_messages(con, cid)
    turns, i = [], 0
    for m in msgs:
        if m["role"] == "user":
            turns.append({"user": m["text"], "assistant": None, "raw": None})
        elif turns:
            turns[-1]["assistant"], turns[-1]["raw"] = m["text"], m["raw"]
    turns.append({"user": new_user_text, "assistant": None, "raw": None})
    last = len(turns) - 1
    for t in turns:
        t["docs"] = []
    for d in docs:  # docs uploaded after the last question ride along with the new one
        turns[min(d["turn"], last)]["docs"].append(d)
    return turns


class FollowupSplitter:
    """Streams answer text through, holding back the follow-up questions after FOLLOWUP_MARK."""

    def __init__(self):
        self.buf, self.tail, self.hit = "", "", False

    def feed(self, s: str) -> str:
        if self.hit:
            self.tail += s
            return ""
        self.buf += s
        i = self.buf.find(FOLLOWUP_MARK)
        if i >= 0:
            out, self.tail, self.buf, self.hit = self.buf[:i], self.buf[i + len(FOLLOWUP_MARK):], "", True
            return out
        keep = 0  # hold back a partial marker at the end of the buffer
        for n in range(min(len(self.buf), len(FOLLOWUP_MARK) - 1), 0, -1):
            if FOLLOWUP_MARK.startswith(self.buf[-n:]):
                keep = n
                break
        out, self.buf = self.buf[:len(self.buf) - keep], self.buf[len(self.buf) - keep:]
        return out

    def flush(self) -> str:
        out, self.buf = ("" if self.hit else self.buf), ""
        return out

    def followups(self):
        lines = [re.sub(r"^\s*(?:[-*•]|\d+[.)])\s*", "", line).strip().strip("*").strip()
                 for line in self.tail.splitlines()]
        return [line for line in lines if 3 < len(line) < 160][:3]


def parse_json_object(text: str) -> dict:
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("no JSON object in reply")
    return json.loads(text[start:end + 1])


def clean_insights(data: dict) -> dict:
    def s(v, n):
        return str(v or "").strip()[:n]
    return {
        "title": s(data.get("title"), 70),
        "type": s(data.get("type"), 40),
        "summary": s(data.get("summary"), 700),
        "topics": [s(t, 40) for t in (data.get("topics") or []) if str(t).strip()][:6],
        "questions": [s(q, 160) for q in (data.get("questions") or []) if str(q).strip()][:4],
    }


# ================================================================ Gemini (free tier)

_gemini_client = None


def gemini_client():
    global _gemini_client
    from google import genai
    if _gemini_client is None:
        if not (os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")):
            raise UserFacingError("No Gemini key found. Get a free one at aistudio.google.com/apikey, "
                                  "put it in the .env file, and restart the server.")
        _gemini_client = genai.Client()
    return _gemini_client


def gemini_doc_parts(docs):
    from google.genai import types
    parts = []
    for d in docs:
        parts.append(types.Part.from_text(text=f'Document: "{d["name"]}"'))
        if d["kind"] in ("pdf", "image"):
            parts.append(types.Part.from_bytes(data=d["raw"], mime_type=d["mime"]))
        else:
            parts.append(types.Part.from_text(text=d["text"] or ""))
    return parts


def gemini_call(fn):
    """Run a Gemini call, turning API errors into friendly messages."""
    from google import genai
    try:
        return fn()
    except genai.errors.APIError as e:
        if e.code == 429:
            raise UserFacingError("Gemini free-tier limit reached. Wait a minute (or until tomorrow for the "
                                  "daily limit), or set GEMINI_MODEL to another model.")
        if e.code in (400, 401, 403) and "key" in str(e.message).lower():
            raise UserFacingError("Gemini rejected the API key. Check GEMINI_API_KEY in .env and restart the server.")
        if e.code == 404:
            raise UserFacingError(f"Gemini model '{GEMINI_MODEL}' not found. Set GEMINI_MODEL to a current model.")
        raise UserFacingError(f"Gemini error {e.code}: {e.message}")


def gemini_stream(turns):
    from google.genai import types
    client = gemini_client()
    contents = []
    for t in turns:
        parts = gemini_doc_parts(t["docs"]) + [types.Part.from_text(text=t["user"])]
        contents.append(types.Content(role="user", parts=parts))
        if t["assistant"] is not None:
            contents.append(types.Content(role="model", parts=[types.Part.from_text(text=t["assistant"])]))
    config = types.GenerateContentConfig(system_instruction=SYSTEM_PROMPT + CITE_HINT)

    def run():
        for chunk in client.models.generate_content_stream(model=GEMINI_MODEL, contents=contents, config=config):
            if chunk.text:
                yield chunk.text

    stream = gemini_call(run)
    while True:
        text = gemini_call(lambda: next(stream, None))
        if text is None:
            break
        yield {"type": "text", "text": text}


def gemini_json(docs, prompt):
    from google.genai import types
    client = gemini_client()
    resp = gemini_call(lambda: client.models.generate_content(
        model=GEMINI_MODEL,
        contents=[types.Content(role="user", parts=gemini_doc_parts(docs) + [types.Part.from_text(text=prompt)])],
        config=types.GenerateContentConfig(response_mime_type="application/json"),
    ))
    return resp.text or ""


# ================================================================ Ollama (local, offline)

def ollama_call(fn):
    import httpx  # ollama's HTTP client; it raises httpx.ConnectError when Ollama isn't running
    import ollama
    try:
        return fn()
    except (ConnectionError, httpx.ConnectError):
        raise UserFacingError("Can't reach Ollama. Install it from ollama.com, make sure it's running, "
                              f"then run: ollama pull {OLLAMA_MODEL}")
    except ollama.ResponseError as e:
        if e.status_code == 404:
            raise UserFacingError(f"Model '{OLLAMA_MODEL}' isn't downloaded yet. Run: ollama pull {OLLAMA_MODEL}")
        raise UserFacingError(f"Ollama error: {e.error}")


def ollama_stream(turns):
    import ollama
    docs = [d for t in turns for d in t["docs"]]
    doc_text = docs_as_text(docs)
    approx_tokens = len(doc_text) // 4
    if len(turns) == 1 and approx_tokens > OLLAMA_CONTEXT * 0.75:
        yield {"type": "text", "text": f"(Heads up: your documents are about {approx_tokens:,} tokens, more than this "
                                       f"local model can hold at once, so it may miss later parts.)\n\n"}
    messages = [{"role": "system", "content": SYSTEM_PROMPT + CITE_HINT + "\n\nThe documents:\n\n" + doc_text}]
    for t in turns:
        messages.append({"role": "user", "content": t["user"]})
        if t["assistant"] is not None:
            messages.append({"role": "assistant", "content": t["assistant"]})
    stream = ollama_call(lambda: ollama.chat(model=OLLAMA_MODEL, messages=messages, stream=True,
                                             options={"num_ctx": OLLAMA_CONTEXT}))
    while True:
        chunk = ollama_call(lambda: next(stream, None))
        if chunk is None:
            break
        if chunk.message.content:
            yield {"type": "text", "text": chunk.message.content}


def ollama_json(docs, prompt):
    import ollama
    resp = ollama_call(lambda: ollama.chat(
        model=OLLAMA_MODEL, format="json", options={"num_ctx": OLLAMA_CONTEXT},
        messages=[{"role": "user", "content": "The documents:\n\n" + docs_as_text(docs) + "\n\n" + prompt}]))
    return resp.message.content or ""


# ================================================================ Claude (paid, exact citations)

def claude_doc_blocks(docs):
    blocks = []
    for d in docs:
        if d["kind"] == "image":
            blocks.append({"type": "image", "source": {"type": "base64", "media_type": d["mime"],
                                                       "data": base64.standard_b64encode(d["raw"]).decode("ascii")}})
            continue
        if d["kind"] == "pdf":
            source = {"type": "base64", "media_type": "application/pdf",
                      "data": base64.standard_b64encode(d["raw"]).decode("ascii")}
        else:
            source = {"type": "text", "media_type": "text/plain", "data": d["text"] or ""}
        blocks.append({"type": "document", "source": source, "title": d["name"], "citations": {"enabled": True}})
    return blocks


def describe_citation(c) -> dict:
    out = {"title": getattr(c, "document_title", None) or "Document", "text": getattr(c, "cited_text", "")}
    if getattr(c, "type", "") == "page_location":
        start, end = c.start_page_number, c.end_page_number - 1  # end is exclusive
        out["where"] = f"p. {start}" if end <= start else f"pp. {start}-{end}"
    return out


def claude_call(fn):
    import anthropic
    try:
        return fn()
    except anthropic.AuthenticationError:
        raise UserFacingError("No valid Anthropic API key. Set ANTHROPIC_API_KEY and restart the server.")
    except anthropic.RateLimitError:
        raise UserFacingError("Rate limited by the API. Wait a moment and try again.")
    except anthropic.APIStatusError as e:
        raise UserFacingError(f"Claude API error {e.status_code}: {e.message}")
    except anthropic.APIConnectionError:
        raise UserFacingError("Couldn't reach the Claude API. Check your internet connection.")


def claude_messages(turns):
    messages = []
    for t in turns:
        messages.append({"role": "user", "content": claude_doc_blocks(t["docs"]) + [{"type": "text", "text": t["user"]}]})
        if t["assistant"] is not None:
            # Replay Claude's own content blocks unchanged (thinking blocks included) so history stays append-only
            messages.append({"role": "assistant", "content": json.loads(t["raw"]) if t["raw"] else t["assistant"]})
    return messages


def claude_stream(turns):
    import anthropic
    client = anthropic.Anthropic()
    citations = []
    events = []

    def run():
        with client.beta.messages.stream(
            model=CLAUDE_MODEL,
            max_tokens=16000,
            system=SYSTEM_PROMPT,
            output_config={"effort": "medium"},
            cache_control={"type": "ephemeral"},  # follow-up questions reuse the cached documents
            messages=claude_messages(turns),
            # If a safety classifier declines, retry server-side on Anthropic's recommended fallback model
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",
        ) as stream:
            for event in stream:
                if event.type != "content_block_delta":
                    continue
                if event.delta.type == "text_delta":
                    yield {"type": "text", "text": event.delta.text}
                elif event.delta.type == "citations_delta":
                    cite = describe_citation(event.delta.citation)
                    if cite not in citations:
                        citations.append(cite)
                        yield {"type": "cite", "n": len(citations)}
            events.append(stream.get_final_message())

    gen = run()
    while True:
        ev = claude_call(lambda: next(gen, None))
        if ev is None:
            break
        yield ev
    final = events[0]
    if final.stop_reason == "refusal":
        raise UserFacingError("Claude declined to answer this one. Try rephrasing the question.")
    if final.stop_reason == "max_tokens":
        yield {"type": "text", "text": "\n\n(Answer was cut off at the length limit.)"}
    yield {"type": "final", "citations": citations,
           "raw": json.dumps([b.to_dict() for b in final.content])}


def claude_json(docs, prompt):
    import anthropic
    resp = claude_call(lambda: anthropic.Anthropic().beta.messages.create(
        model=CLAUDE_MODEL, max_tokens=4000, output_config={"effort": "low"},
        messages=[{"role": "user", "content": claude_doc_blocks(docs) + [{"type": "text", "text": prompt}]}],
        betas=["server-side-fallback-2026-07-01"], fallbacks="default",
    ))
    return "".join(b.text for b in resp.content if b.type == "text")


PROVIDERS = {
    "gemini": (gemini_stream, gemini_json),
    "ollama": (ollama_stream, ollama_json),
    "claude": (claude_stream, claude_json),
}


def provider():
    if PROVIDER not in PROVIDERS:
        raise UserFacingError(f"Unknown PROVIDER '{PROVIDER}'. Use gemini, ollama or claude.")
    return PROVIDERS[PROVIDER]


# ================================================================ routes

def error(msg, code=400):
    return jsonify({"error": msg}), code


@app.get("/")
def index():
    if AUTH_ENABLED and not g.user:
        return redirect("/login")
    return send_from_directory("static", "index.html")


@app.get("/login")
def login_page():
    """Separate sign-up / log-in page; signed-in visitors go straight to the app."""
    if not AUTH_ENABLED or g.user:
        return redirect("/")
    return send_from_directory("static", "login.html")


@app.get("/privacy")
def privacy():
    return send_from_directory("static", "privacy.html")


@app.get("/api/config")
def config():
    remaining = None
    if DAILY_LIMIT:
        with db() as con:
            remaining = max(0, DAILY_LIMIT - usage_today(con, "ip:" + client_ip()))
    return jsonify({"provider": PROVIDER, "label": PROVIDER_LABEL, "images": PROVIDER != "ollama",
                    "citations": PROVIDER == "claude", "public": PUBLIC, "limit": DAILY_LIMIT,
                    "remaining": remaining, "max_mb": MAX_UPLOAD_BYTES // (1024 * 1024),
                    "auth": AUTH_ENABLED, "google_client_id": GOOGLE_CLIENT_ID or None, "user": g.user})


@app.post("/api/auth/google")
def google_login():
    """Verify the ID token from the 'Sign in with Google' button and start a session."""
    if not AUTH_ENABLED:
        return error("Google sign-in isn't set up on this server.", 404)
    credential = (request.get_json(silent=True) or {}).get("credential", "")
    try:
        from google.auth.transport import requests as google_requests
        from google.oauth2 import id_token
        info = id_token.verify_oauth2_token(credential, google_requests.Request(), GOOGLE_CLIENT_ID,
                                            clock_skew_in_seconds=10)
    except Exception:
        return error("Google sign-in failed. Please try again.", 401)
    user = {"id": "g:" + info["sub"], "email": info.get("email", ""),
            "name": info.get("name") or info.get("email", "").split("@")[0], "picture": info.get("picture", "")}
    now = time.time()
    with db() as con:
        con.execute("INSERT INTO users (id, email, name, picture, created, last_login) VALUES (?, ?, ?, ?, ?, ?) "
                    "ON CONFLICT(id) DO UPDATE SET email = excluded.email, name = excluded.name, "
                    "picture = excluded.picture, last_login = excluded.last_login",
                    (user["id"], user["email"], user["name"], user["picture"], now, now))
        # Chats made in this browser before signing in move into the account
        con.execute("UPDATE chats SET owner = ? WHERE owner = ?", (user["id"], g.guest))
    session.clear()
    session["user"] = user
    session.permanent = True
    return jsonify({"user": user})


@app.post("/api/auth/logout")
def logout():
    session.clear()
    return jsonify({"ok": True})


@app.get("/api/chats")
def list_chats():
    with db() as con:
        if not PUBLIC and not AUTH_ENABLED:  # on your own computer, chats saved before per-visitor privacy are yours
            con.execute("UPDATE chats SET owner = ? WHERE owner IS NULL", (g.owner,))
        rows = con.execute("SELECT * FROM chats WHERE owner = ? ORDER BY updated DESC LIMIT 200", (g.owner,)).fetchall()
        return jsonify([chat_summary(con, r) for r in rows])


@app.post("/api/chats")
def create_chat():
    cid, now = uuid.uuid4().hex[:12], time.time()
    with db() as con:
        con.execute("INSERT INTO chats (id, owner, created, updated) VALUES (?, ?, ?, ?)", (cid, g.owner, now, now))
        return jsonify(chat_summary(con, get_chat(con, cid)))


@app.get("/api/chats/<cid>")
def open_chat(cid):
    with db() as con:
        chat = get_chat(con, cid)
        if not chat:
            return error("Chat not found.", 404)
        msgs = [{"id": m["id"], "role": m["role"], "display": m["display"], "meta": json.loads(m["meta"] or "{}")}
                for m in load_messages(con, cid)]
        return jsonify({**chat_summary(con, chat),
                        "insights": json.loads(chat["insights"]) if chat["insights"] else None,
                        "docs": [doc_public(d) for d in load_docs(con, cid)], "messages": msgs})


@app.patch("/api/chats/<cid>")
def rename_chat(cid):
    title = ((request.get_json(silent=True) or {}).get("title") or "").strip()[:80]
    with db() as con:
        if not get_chat(con, cid):
            return error("Chat not found.", 404)
        con.execute("UPDATE chats SET title = ? WHERE id = ?", (title, cid))
    return jsonify({"ok": True})


@app.delete("/api/chats/<cid>")
def delete_chat(cid):
    with db() as con:
        if not get_chat(con, cid):
            return error("Chat not found.", 404)
        con.execute("DELETE FROM messages WHERE chat_id = ?", (cid,))
        con.execute("DELETE FROM docs WHERE chat_id = ?", (cid,))
        con.execute("DELETE FROM chats WHERE id = ?", (cid,))
    return jsonify({"ok": True})


@app.post("/api/chats/<cid>/upload")
def upload(cid):
    files = request.files.getlist("files")
    if not files:
        return error("No files received.")
    added, errors = [], []
    with db() as con:
        if not get_chat(con, cid):
            return error("Chat not found.", 404)
        turn = con.execute("SELECT COUNT(*) FROM messages WHERE chat_id = ? AND role = 'user'", (cid,)).fetchone()[0]
        total = con.execute("SELECT COALESCE(SUM(size), 0) FROM docs WHERE chat_id = ?", (cid,)).fetchone()[0]
        for f in files:
            name = os.path.basename(f.filename or "untitled")
            raw = f.read()
            if total + len(raw) > MAX_UPLOAD_BYTES:
                errors.append(f"{name}: skipped, this chat's files would exceed {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")
                continue
            try:
                doc = ingest(name, raw)
            except UserFacingError as e:
                errors.append(str(e))
                continue
            con.execute("INSERT INTO docs (id, chat_id, name, kind, mime, size, raw, text, meta, turn, created) "
                        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        (uuid.uuid4().hex[:12], cid, name, doc["kind"], doc["mime"], len(raw), doc["raw"],
                         doc["text"], json.dumps(doc["meta"]), turn, time.time()))
            total += len(raw)
            added.append(name)
        if added:
            touch(con, cid)
        return jsonify({"added": added, "errors": errors, "docs": [doc_public(d) for d in load_docs(con, cid)]})


@app.delete("/api/chats/<cid>/docs/<did>")
def delete_doc(cid, did):
    with db() as con:
        if not get_chat(con, cid):
            return error("Chat not found.", 404)
        sent = con.execute("SELECT turn < (SELECT COUNT(*) FROM messages WHERE chat_id = ? AND role = 'user') "
                           "FROM docs WHERE id = ? AND chat_id = ?", (cid, did, cid)).fetchone()
        if sent is None:
            return error("Document not found.", 404)
        # Claude replays its earlier turns exactly, so a document it has already read can't be pulled out.
        # Gemini and Ollama rebuild the conversation each time, so removing it is fine.
        if sent[0] and PROVIDER == "claude":
            return error("This document is already part of the conversation. Start a new chat to drop it.")
        con.execute("DELETE FROM docs WHERE id = ? AND chat_id = ?", (did, cid))
        return jsonify({"docs": [doc_public(d) for d in load_docs(con, cid)]})


@app.post("/api/chats/<cid>/insights")
def insights(cid):
    with db() as con:
        chat = get_chat(con, cid)
        if not chat:
            return error("Chat not found.", 404)
        docs = load_docs(con, cid, with_data=True)
    if not docs:
        return error("No documents yet.")
    try:
        take_quota()
        text = provider()[1](docs, INSIGHTS_PROMPT)
        data = clean_insights(parse_json_object(text))
    except UserFacingError as e:
        return error(str(e), 429 if isinstance(e, QuotaError) else 502)
    except (ValueError, json.JSONDecodeError):
        return error("Couldn't build an overview for these documents.", 502)
    with db() as con:
        con.execute("UPDATE chats SET insights = ? WHERE id = ?", (json.dumps(data), cid))
        if not chat["title"] and data["title"]:
            con.execute("UPDATE chats SET title = ? WHERE id = ?", (data["title"], cid))
    return jsonify(data)


@app.post("/api/chats/<cid>/undo")
def undo_last(cid):
    """Remove the last question and answer (used by Retry). Returns what was asked."""
    with db() as con:
        if not get_chat(con, cid):
            return error("Chat not found.", 404)
        msgs = load_messages(con, cid)
        if len(msgs) < 2 or msgs[-1]["role"] != "assistant" or msgs[-2]["role"] != "user":
            return error("Nothing to retry.")
        con.execute("DELETE FROM messages WHERE id IN (?, ?)", (msgs[-1]["id"], msgs[-2]["id"]))
        meta = json.loads(msgs[-2]["meta"] or "{}")
        return jsonify({"question": meta.get("prompt") or msgs[-2]["display"], "display": msgs[-2]["display"],
                        "style": meta.get("style", "balanced")})


@app.post("/api/chats/<cid>/ask")
def ask(cid):
    body = request.get_json(silent=True) or {}
    prompt = (body.get("question") or "").strip()
    display = (body.get("display") or prompt).strip()
    style = body.get("style") or "balanced"
    if not prompt:
        return error("Empty question.")
    with db() as con:
        chat = get_chat(con, cid)
        if not chat:
            return error("Chat not found.", 404)
        docs = load_docs(con, cid, with_data=True)
        if not docs:
            return error("Add a document first.")
        sent = prompt + (f"\n\n(Answer style: {STYLE_HINTS[style]})" if style in STYLE_HINTS else "")
        turns = build_turns(con, cid, docs, sent)
    try:
        stream_fn = provider()[0]
        take_quota()
    except UserFacingError as e:
        return error(str(e), 429 if isinstance(e, QuotaError) else 500)

    def save(answer, shown, raw, meta):
        with db() as con:
            now = time.time()
            con.execute("INSERT INTO messages (chat_id, role, text, display, meta, created) VALUES (?, 'user', ?, ?, ?, ?)",
                        (cid, sent, display, json.dumps({"prompt": prompt, "style": style}), now))
            con.execute("INSERT INTO messages (chat_id, role, text, display, raw, meta, created) "
                        "VALUES (?, 'assistant', ?, ?, ?, ?, ?)", (cid, answer, shown, raw, json.dumps(meta), now))
            if not chat["title"]:
                con.execute("UPDATE chats SET title = ? WHERE id = ?", (display[:60], cid))
            touch(con, cid)

    def generate():
        splitter = FollowupSplitter()
        answer, shown = [], []  # shown keeps citation markers for redisplay
        state = {"saved": False, "failed": False, "raw": None, "citations": []}
        try:
            for ev in stream_fn(turns):
                if ev["type"] == "text":
                    out = splitter.feed(ev["text"])
                    if out:
                        answer.append(out)
                        shown.append(out)
                        yield sse({"type": "text", "text": out})
                elif ev["type"] == "cite":
                    shown.append(f"\u0001{ev['n']}\u0001")
                    yield sse(ev)
                elif ev["type"] == "final":
                    state["raw"], state["citations"] = ev.get("raw"), ev.get("citations", [])
            tail = splitter.flush()
            if tail:
                answer.append(tail)
                shown.append(tail)
                yield sse({"type": "text", "text": tail})
            meta = {"citations": state["citations"], "followups": splitter.followups()}
            save("".join(answer).strip(), "".join(shown).strip(), state["raw"], meta)
            state["saved"] = True
            yield sse({"type": "done", **meta})
        except UserFacingError as e:
            state["failed"] = True
            yield sse({"type": "error", "text": str(e)})
        except Exception as e:
            state["failed"] = True
            app.logger.exception("chat failed")
            yield sse({"type": "error", "text": f"Unexpected error: {e}"})
        finally:
            # The browser pressed Stop: keep the partial answer so the chat still makes sense
            if not state["saved"] and not state["failed"] and "".join(answer).strip():
                partial = "".join(answer).strip() + "\n\n(Stopped.)"
                save(partial, "".join(shown).strip() + "\n\n_(Stopped.)_", None, {"citations": [], "followups": []})

    return Response(stream_with_context(generate()), mimetype="text/event-stream",
                    headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/api/chats/<cid>/export")
def export_chat(cid):
    with db() as con:
        chat = get_chat(con, cid)
        if not chat:
            return error("Chat not found.", 404)
        docs, msgs = load_docs(con, cid), load_messages(con, cid)
    title = chat["title"] or "DocChat conversation"
    lines = [f"# {title}", "", f"Exported from DocChat on {time.strftime('%Y-%m-%d %H:%M')}", ""]
    if docs:
        lines += ["**Documents:** " + ", ".join(d["name"] for d in docs), ""]
    for m in msgs:
        text = re.sub(r"\u0001(\d+)\u0001", r"[\1]", m["display"])
        lines += [f"## {'You' if m['role'] == 'user' else 'DocChat'}", "", text, ""]
        for i, c in enumerate(json.loads(m["meta"] or "{}").get("citations") or [], 1):
            lines.append(f"> [{i}] {c['title']}{', ' + c['where'] if c.get('where') else ''}: \"{c['text'].strip()}\"")
    safe = re.sub(r"[^\w\- ]", "", title).strip().replace(" ", "-")[:50] or "docchat"
    return Response("\n".join(lines), mimetype="text/markdown",
                    headers={"Content-Disposition": f'attachment; filename="{safe}.md"'})


init_db()

if __name__ == "__main__":
    print(f"DocChat using: {PROVIDER_LABEL}".encode("ascii", "replace").decode())
    app.run(host="127.0.0.1", port=5000, debug=False, threaded=True)
